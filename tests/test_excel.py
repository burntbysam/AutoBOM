from __future__ import annotations

from decimal import Decimal

from openpyxl import load_workbook

from autobom.core.excel import (
    COLUMNS,
    SHEET_ORDER,
    build_summary,
    select_sheets,
    tally,
    write_workbook,
)
from autobom.core.models import Part


def part(number, thickness='1/8"', fits="T", quantity="1", size="60x120") -> Part:
    return Part(
        quantity=Decimal(quantity),
        part_number=number,
        thickness_label=thickness,
        size=size,
        fits_trumpf=fits,
    )


PARTS = [
    part("8701-1101-1", fits="F", size="72x120"),
    part("8701-1101-2"),
    part("8701-1101-3", thickness='3/16"'),
    part("8701-1101-4", thickness="OTHER"),
    part("8701-1101-5", thickness="OTHER", fits="F", size="72x144"),
    part("8701-1101-6", thickness='3/16"', fits="F", size="72x120"),
]


class TestSheetRouting:
    def test_eighth_and_t_only(self):
        sheets = select_sheets(PARTS)
        assert [p.part_number for p in sheets["1-8"]] == ["8701-1101-2"]

    def test_three_sixteenth_and_t_only(self):
        sheets = select_sheets(PARTS)
        assert [p.part_number for p in sheets["3-16"]] == ["8701-1101-3"]

    def test_f_parts_includes_all_thicknesses(self):
        sheets = select_sheets(PARTS)
        assert [p.part_number for p in sheets["F Parts"]] == [
            "8701-1101-1",
            "8701-1101-5",
            "8701-1101-6",
        ]

    def test_other_includes_both_fits(self):
        sheets = select_sheets(PARTS)
        assert [p.part_number for p in sheets["Other"]] == ["8701-1101-4", "8701-1101-5"]

    def test_all_contains_everything(self):
        assert len(select_sheets(PARTS)["All"]) == len(PARTS)

    def test_other_never_lands_on_a_thickness_sheet(self):
        sheets = select_sheets(PARTS)
        for name in ("1-8", "3-16"):
            assert all(p.thickness_label != "OTHER" for p in sheets[name])

    def test_f_never_lands_on_a_thickness_sheet(self):
        sheets = select_sheets(PARTS)
        for name in ("1-8", "3-16"):
            assert all(p.fits_trumpf == "T" for p in sheets[name])


class TestWorkbook:
    def test_exactly_five_sheets_in_order(self, tmp_path):
        path = write_workbook(PARTS, tmp_path / "out.xlsx")
        workbook = load_workbook(path)
        assert workbook.sheetnames == list(SHEET_ORDER)

    def test_headers_on_every_sheet(self, tmp_path):
        workbook = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))
        for name in SHEET_ORDER:
            header = [cell.value for cell in workbook[name][1]]
            assert header == list(COLUMNS)

    def test_no_extra_columns(self, tmp_path):
        workbook = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))
        for name in SHEET_ORDER:
            assert workbook[name].max_column == len(COLUMNS)

    def test_alignment_matches_spec(self, tmp_path):
        workbook = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))
        sheet = workbook["All"]
        expected = ["left", "center", "center", "center", "left"]
        for row in sheet.iter_rows(min_row=2, max_row=2):
            assert [cell.alignment.horizontal for cell in row] == expected

    def test_whole_quantities_written_as_integers(self, tmp_path):
        workbook = load_workbook(
            write_workbook([part("8701-1101-1", quantity="4")], tmp_path / "out.xlsx")
        )
        value = workbook["All"].cell(row=2, column=1).value
        assert value == 4
        assert isinstance(value, int)

    def test_empty_sheets_still_exist_with_headers(self, tmp_path):
        workbook = load_workbook(write_workbook([], tmp_path / "out.xlsx"))
        assert workbook.sheetnames == list(SHEET_ORDER)
        for name in SHEET_ORDER:
            # Header, a blank separator, then the totals block.
            assert [cell.value for cell in workbook[name][1]] == list(COLUMNS)
            assert workbook[name][2][0].value is None

    def test_empty_sheet_totals_are_zero(self, tmp_path):
        workbook = load_workbook(write_workbook([], tmp_path / "out.xlsx"))
        row = workbook["1-8"][3]
        assert row[0].value == 0
        assert row[1].value == "TOTAL PIECES"
        assert row[2].value == 0


class TestTally:
    def test_counts_line_items_and_pieces(self):
        assert tally(PARTS) == (6, 6)

    def test_pieces_sums_the_quantity_column(self):
        parts = [part("A", quantity="3"), part("B", quantity="4")]
        assert tally(parts) == (2, 7)

    def test_empty(self):
        assert tally([]) == (0, 0)

    def test_whole_totals_stay_integers(self):
        line_items, pieces = tally([part("A", quantity="2"), part("B", quantity="2")])
        assert isinstance(pieces, int)


class TestSummary:
    def test_groups_by_thickness_regardless_of_trumpf_fit(self):
        summary = dict((label, (items, pieces)) for label, items, pieces in build_summary(PARTS))
        # 1/8": one T part and one F part.
        assert summary['1/8"'] == (2, 2)
        # 3/16": one T part and one F part.
        assert summary['3/16"'] == (2, 2)
        assert summary["OTHER thickness"] == (2, 2)

    def test_combined_row_is_the_two_stock_thicknesses(self):
        summary = dict((label, (items, pieces)) for label, items, pieces in build_summary(PARTS))
        assert summary['1/8" + 3/16" total'] == (4, 4)

    def test_rows_add_up_to_the_grand_total(self):
        summary = dict((label, (items, pieces)) for label, items, pieces in build_summary(PARTS))
        parts = summary['1/8"'][0] + summary['3/16"'][0] + summary["OTHER thickness"][0]
        pieces = summary['1/8"'][1] + summary['3/16"'][1] + summary["OTHER thickness"][1]
        assert (parts, pieces) == summary["Grand total"]

    def test_quantities_are_summed_not_counted(self):
        parts = [
            part("A", quantity="5"),
            part("B", thickness='3/16"', quantity="7"),
            part("C", thickness="OTHER", quantity="2"),
        ]
        summary = dict((label, (items, pieces)) for label, items, pieces in build_summary(parts))
        assert summary['1/8"'] == (1, 5)
        assert summary['1/8" + 3/16" total'] == (2, 12)
        assert summary["Grand total"] == (3, 14)


class TestWorkbookTotals:
    def test_summary_block_is_on_the_all_sheet(self, tmp_path):
        sheet = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))["All"]
        labels = [row[1].value for row in sheet.iter_rows(min_row=2)]
        assert "SUMMARY" in labels
        for expected in ('1/8"', '3/16"', '1/8" + 3/16" total', "OTHER thickness", "Grand total"):
            assert expected in labels

    def test_category_sheets_get_a_totals_line(self, tmp_path):
        workbook = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))
        for name in ("1-8", "3-16", "F Parts", "Other"):
            labels = [row[1].value for row in workbook[name].iter_rows(min_row=2)]
            assert "TOTAL PIECES" in labels

    def test_sheet_total_matches_its_rows(self, tmp_path):
        workbook = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))
        for name, rows in select_sheets(PARTS).items():
            if name == "All":
                continue
            sheet = workbook[name]
            total_row = next(
                row for row in sheet.iter_rows(min_row=2) if row[1].value == "TOTAL PIECES"
            )
            assert total_row[0].value == tally(rows)[1]
            assert total_row[2].value == tally(rows)[0]

    def test_summary_is_separated_from_the_data_by_a_blank_row(self, tmp_path):
        sheet = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))["All"]
        # Data occupies rows 2..7 for six parts; row 8 must be blank.
        assert all(cell.value is None for cell in sheet[len(PARTS) + 2])

    def test_data_rows_are_unaffected(self, tmp_path):
        sheet = load_workbook(write_workbook(PARTS, tmp_path / "out.xlsx"))["All"]
        rows = [
            tuple(row)
            for row in sheet.iter_rows(
                min_row=2, max_row=len(PARTS) + 1, values_only=True
            )
        ]
        assert [row[1] for row in rows] == [p.part_number for p in PARTS]
