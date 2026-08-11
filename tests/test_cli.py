from __future__ import annotations

from openpyxl import load_workbook

from autobom.cli import main, split_inputs

from .conftest import write_csv

EIGHTH = "SHEET,AL,SMOOTH,3003,.125,60x120"


def make_job(tmp_path, extra_il_rows=()):
    write_csv(tmp_path, "8701-01101-I.csv", [f"1|1|{EIGHTH}|X|A|"])
    write_csv(
        tmp_path,
        "IL-8701-011.csv",
        ["1|2|BUS SECTION|8701-01101-I|", *extra_il_rows],
    )
    return tmp_path


class TestSplitInputs:
    def test_il_prefix_routes_to_il_list(self, tmp_path):
        boms, ils = split_inputs(
            [tmp_path / "8701-01101-I.csv", tmp_path / "IL-8701-011.csv"]
        )
        assert [p.name for p in boms] == ["8701-01101-I.csv"]
        assert [p.name for p in ils] == ["IL-8701-011.csv"]


class TestMain:
    def test_writes_workbook_from_a_directory(self, tmp_path, capsys):
        job = make_job(tmp_path)
        out = tmp_path / "out.xlsx"
        assert main([str(job), "-o", str(out)]) == 0
        workbook = load_workbook(out)
        assert workbook["All"].cell(row=2, column=1).value == 2
        assert workbook["All"].cell(row=2, column=2).value == "8701-1101-1"
        assert "Wrote" in capsys.readouterr().out

    def test_flags_block_the_write(self, tmp_path, capsys):
        job = make_job(tmp_path, extra_il_rows=["2|1|SPLICE|JB-2724-06|"])
        out = tmp_path / "out.xlsx"
        assert main([str(job), "-o", str(out)]) == 1
        assert not out.exists()
        output = capsys.readouterr().out
        assert "FLAGS — REVIEW REQUIRED" in output
        assert "JB-2724-06" in output

    def test_ignore_flags_writes_anyway(self, tmp_path):
        job = make_job(tmp_path, extra_il_rows=["2|1|SPLICE|JB-2724-06|"])
        out = tmp_path / "out.xlsx"
        assert main([str(job), "-o", str(out), "--ignore-flags"]) == 0
        # The unmatched IL assembly must not appear as a fabricated part row.
        numbers = [
            row[1] for row in load_workbook(out)["All"].iter_rows(min_row=2, values_only=True)
        ]
        assert numbers == ["8701-1101-1"]

    def test_missing_file_is_an_error(self, tmp_path, capsys):
        assert main([str(tmp_path / "nope.csv"), "-o", str(tmp_path / "o.xlsx")]) == 2

    def test_no_il_is_an_error(self, tmp_path, capsys):
        write_csv(tmp_path, "8701-01101-I.csv", [f"1|1|{EIGHTH}|X|A|"])
        assert main([str(tmp_path), "-o", str(tmp_path / "o.xlsx")]) == 2
        assert "no IL CSVs found" in capsys.readouterr().err
