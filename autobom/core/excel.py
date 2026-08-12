"""PHASE 6 workbook writer: five sheets, identical columns and formatting."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .classify import quantity_value
from .models import Part

COLUMNS = ("Qty", "Part #", "Thickness", "Size", "Fits Trumpf")
# PHASE 5 alignment, one entry per column above.
ALIGNMENTS = ("left", "center", "center", "center", "left")
COLUMN_WIDTHS = (10, 22, 12, 16, 13)

SHEET_ORDER = ("1-8", "3-16", "F Parts", "Other", "All")

_BOLD = Font(bold=True)


def select_sheets(parts: list[Part]) -> dict[str, list[Part]]:
    """Route each part to its sheets per the SPEC's critical reminders.

    A part reaches "1-8" or "3-16" only when it is both that thickness and a
    Trumpf fit; anything F goes to "F Parts" and anything OTHER goes to
    "Other". Every part appears on "All".
    """
    return {
        "1-8": [part for part in parts if part.is_eighth and part.fits],
        "3-16": [part for part in parts if part.is_three_sixteenth and part.fits],
        "F Parts": [part for part in parts if not part.fits],
        "Other": [part for part in parts if part.is_other],
        "All": list(parts),
    }


def tally(parts: list[Part]) -> tuple[int, int | float]:
    """``(line items, pieces)`` -- distinct part numbers, and how many to cut."""
    pieces = sum((part.quantity for part in parts), Decimal(0))
    return len(parts), quantity_value(pieces)


def build_summary(parts: list[Part]) -> list[tuple[str, int, int | float]]:
    """Totals by thickness, counting a part whether or not it fits the Trumpf.

    Grouping by thickness rather than by sheet means the rows add up: 1/8" plus
    3/16" plus OTHER is the grand total, with no part counted twice and none
    left out. The per-sheet totals cover the fits-Trumpf split.
    """
    eighth = [part for part in parts if part.is_eighth]
    three_sixteenth = [part for part in parts if part.is_three_sixteenth]
    other = [part for part in parts if part.is_other]
    return [
        ('1/8"', *tally(eighth)),
        ('3/16"', *tally(three_sixteenth)),
        ('1/8" + 3/16" total', *tally(eighth + three_sixteenth)),
        ("OTHER thickness", *tally(other)),
        ("Grand total", *tally(parts)),
    ]


def _write_header(sheet: Worksheet) -> None:
    sheet.append(list(COLUMNS))
    for index, (alignment, width) in enumerate(zip(ALIGNMENTS, COLUMN_WIDTHS), start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        header = sheet.cell(row=1, column=index)
        header.font = _BOLD
        header.alignment = Alignment(horizontal=alignment, vertical="center")


def _write_rows(sheet: Worksheet, parts: list[Part]) -> None:
    for part in parts:
        sheet.append(
            [
                quantity_value(part.quantity),
                part.part_number,
                part.thickness_label,
                part.size,
                part.fits_trumpf,
            ]
        )
    # Applied before any totals block is appended, so the totals keep their own
    # formatting.
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(COLUMNS)):
        for cell, alignment in zip(row, ALIGNMENTS):
            cell.alignment = Alignment(horizontal=alignment, vertical="center")


def _put(sheet: Worksheet, row: int, column: int, value, bold: bool = False):
    cell = sheet.cell(row=row, column=column, value=value)
    if bold:
        cell.font = _BOLD
    cell.alignment = Alignment(
        horizontal=ALIGNMENTS[column - 1] if column <= len(ALIGNMENTS) else "left",
        vertical="center",
    )
    return cell


def _write_sheet_total(sheet: Worksheet, parts: list[Part]) -> None:
    """One totals line, with the piece count under the Qty column it sums."""
    line_items, pieces = tally(parts)
    row = sheet.max_row + 2
    _put(sheet, row, 1, pieces, bold=True)
    _put(sheet, row, 2, "TOTAL PIECES", bold=True)
    _put(sheet, row, 3, line_items, bold=True)
    _put(sheet, row, 4, "LINE ITEMS", bold=True)


def _write_summary(sheet: Worksheet, parts: list[Part]) -> None:
    """The cross-category tally, written below the data on the All sheet."""
    row = sheet.max_row + 2
    _put(sheet, row, 2, "SUMMARY", bold=True)

    row += 1
    _put(sheet, row, 2, "Category", bold=True)
    _put(sheet, row, 3, "Line items", bold=True)
    _put(sheet, row, 4, "Pieces", bold=True)

    for label, line_items, pieces in build_summary(parts):
        row += 1
        emphasis = label.endswith("total")
        _put(sheet, row, 2, label, bold=emphasis)
        _put(sheet, row, 3, line_items, bold=emphasis)
        _put(sheet, row, 4, pieces, bold=emphasis)


def _write_sheet(sheet: Worksheet, parts: list[Part], with_summary: bool) -> None:
    _write_header(sheet)
    _write_rows(sheet, parts)
    if with_summary:
        _write_summary(sheet, parts)
    else:
        _write_sheet_total(sheet, parts)
    sheet.freeze_panes = "A2"


def build_workbook(parts: list[Part]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = select_sheets(parts)
    for name in SHEET_ORDER:
        _write_sheet(workbook.create_sheet(title=name), sheets[name], name == "All")
    return workbook


def write_workbook(parts: list[Part], destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(parts).save(destination)
    return destination
